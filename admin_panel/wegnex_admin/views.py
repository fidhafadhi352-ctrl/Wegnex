from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth import login, logout
from django.contrib.auth.models import User
from django.contrib import messages
from django.views.decorators.cache import never_cache
from user.account.models import Category


@never_cache
def admin_login(request):

    if request.user.is_authenticated and (request.user.is_staff or request.user.is_superuser):
        return redirect("admin_dashboard")

    if request.method == "POST":
        identifier = request.POST.get("email", "").strip() or request.POST.get("username", "").strip()
        password = request.POST.get("password", "")

        if not identifier:
            messages.error(request, "Email address or username is required.")
            return render(request, "admin_login.html")


        if not password:
            messages.error(request, "Password is required.")
            return render(request, "admin_login.html", {"identifier": identifier})

        user = User.objects.filter(email__iexact=identifier).first()
        if not user:
            user = User.objects.filter(username__iexact=identifier).first()

        if not user:
            messages.error(request, "Invalid administrator credentials.")
            return render(request, "admin_login.html", {"identifier": identifier})

        if not user.is_staff and not user.is_superuser:
            messages.error(request, "You are not authorized to access the administrator panel.")
            return render(request, "admin_login.html", {"identifier": identifier})
        
        if not user.is_active:
            messages.error(request, "This administrator account has been deactivated.")
            return render(request, "admin_login.html", {"identifier": identifier})

        if not user.check_password(password):
            messages.error(request, "Invalid administrator credentials.")
            return render(request, "admin_login.html", {"identifier": identifier})

        login(request, user, backend='django.contrib.auth.backends.ModelBackend')
        messages.success(request, f"Welcome back, {user.first_name or user.username}!")
        return redirect("admin_dashboard")

    return render(request, "admin_login.html")


@never_cache
def admin_logout(request):
    logout(request)
    messages.success(request, "You have been logged out of the administrator portal.")
    return redirect("admin_login")


@never_cache
def admin_dashboard(request):

    if not request.user.is_authenticated or (not request.user.is_staff and not request.user.is_superuser):
        messages.error(request, "Please log in with administrator privileges.")
        return redirect("admin_login")

    if not request.user.is_active:
        return redirect("admin_login")

    import datetime
    import calendar
    from django.utils import timezone
    from django.db.models import Sum, Count, Q, F
    from user.checkout.models import Order, OrderItem
    from admin_panel.product_management.models import Product

    # ---- Customer stats ----
    customers_qs = User.objects.filter(is_staff=False, is_superuser=False)
    total_customers = customers_qs.count()
    active_users = customers_qs.filter(is_active=True).count()
    blocked_users = customers_qs.filter(is_active=False).count()

    # ---- Category stats ----
    categories_qs = Category.objects.filter(is_deleted=False).order_by('-created_at')
    total_categories = categories_qs.count()
    active_categories = categories_qs.filter(is_active=True).count()

    # ---- Order / Revenue stats ----
    all_orders = Order.objects.all()
    valid_orders = all_orders.exclude(order_status='CANCELLED')
    total_revenue = valid_orders.aggregate(s=Sum('total_amount'))['s'] or 0
    total_orders_count = all_orders.count()
    pending_orders_count = all_orders.filter(order_status__in=['PENDING', 'CONFIRMED', 'PROCESSING']).count()

    # ---- Top Selling Products (by total quantity sold) ----
    top_products_qs = (
        OrderItem.objects
        .exclude(order__order_status='CANCELLED')
        .values('product_name', 'product__category__name', 'price', 'product__is_active')
        .annotate(total_sold=Sum('quantity'), revenue=Sum('item_total'))
        .order_by('-total_sold')[:5]
    )
    top_products = list(top_products_qs)

    # ---- Top Categories (by units sold) ----
    top_categories_qs = (
        OrderItem.objects
        .exclude(order__order_status='CANCELLED')
        .filter(product__category__isnull=False)
        .values('product__category__name')
        .annotate(total_sold=Sum('quantity'))
        .order_by('-total_sold')[:6]
    )
    top_categories_raw = list(top_categories_qs)
    max_cat_sold = top_categories_raw[0]['total_sold'] if top_categories_raw else 1
    top_categories = []
    for c in top_categories_raw:
        pct = round((c['total_sold'] / max_cat_sold) * 100) if max_cat_sold else 0
        top_categories.append({
            'name': c['product__category__name'],
            'sold': c['total_sold'],
            'pct': pct,
        })

    # ---- Revenue Overview - last 12 months (monthly) ----
    today = timezone.localdate()
    revenue_labels = []
    revenue_data = []
    cur_year = today.year
    cur_month = today.month
    for i in range(11, -1, -1):
        # Calculate month going back i months
        total_months = cur_year * 12 + (cur_month - 1) - i
        yr = total_months // 12
        mo = total_months % 12 + 1
        label = calendar.month_abbr[mo]
        rev = float(
            valid_orders
            .filter(created_at__year=yr, created_at__month=mo)
            .aggregate(s=Sum('total_amount'))['s'] or 0
        )
        revenue_labels.append(label)
        revenue_data.append(rev)

    max_revenue = max(revenue_data) if any(revenue_data) else 1

    # ---- Recent Activity (last 8 orders) ----
    recent_orders = (
        all_orders
        .select_related('user')
        .prefetch_related('items')
        .order_by('-created_at')[:8]
    )
    recent_activity = []
    for o in recent_orders:
        first_item = o.items.first()
        product_name = first_item.product_name if first_item else 'Order'
        time_diff = timezone.now() - o.created_at
        minutes = int(time_diff.total_seconds() / 60)
        if minutes < 60:
            time_str = f"{minutes} min{'s' if minutes != 1 else ''} ago"
        elif minutes < 1440:
            hours = minutes // 60
            time_str = f"{hours} hour{'s' if hours != 1 else ''} ago"
        else:
            days = minutes // 1440
            time_str = f"{days} day{'s' if days != 1 else ''} ago"

        status_display = o.get_order_status_display()
        recent_activity.append({
            'title': f"{product_name}",
            'order_number': o.order_number,
            'time': time_str,
            'status': o.order_status,
            'status_display': status_display,
            'amount': o.total_amount,
            'customer': o.shipping_full_name or o.user.get_full_name() or o.user.username,
        })

    return render(request, "admin_dashboard.html", {
        "admin_user": request.user,
        "total_customers": total_customers,
        "active_users": active_users,
        "blocked_users": blocked_users,
        "total_categories": total_categories,
        "active_categories": active_categories,
        "total_revenue": total_revenue,
        "total_orders_count": total_orders_count,
        "pending_orders_count": pending_orders_count,
        "top_products": top_products,
        "top_categories": top_categories,
        "revenue_labels": revenue_labels,
        "revenue_data": revenue_data,
        "max_revenue": max_revenue if max_revenue else 1,
        "recent_activity": recent_activity,
    })

import csv
import io
import datetime
from django.utils import timezone
from django.db.models import Sum, Count, Q
from django.http import HttpResponse, JsonResponse
from user.checkout.models import Order, OrderItem


@never_cache
def sales_report(request):
    if not request.user.is_authenticated or (not request.user.is_staff and not request.user.is_superuser):
        messages.error(request, "Please log in with administrator privileges.")
        return redirect("admin_login")

    today = timezone.localdate()
    filter_type = request.GET.get("filter_type", "all").strip().lower()
    status_filter = request.GET.get("status", "all").strip()
    export_format = request.GET.get("export", "").strip().lower()

    start_date = None
    end_date = None
    specific_date = None
    date_error = None

    if filter_type == "daily":
        specific_date_str = request.GET.get("date", "").strip()
        if specific_date_str:
            try:
                specific_date = datetime.datetime.strptime(specific_date_str, "%Y-%m-%d").date()
            except ValueError:
                specific_date = today
                date_error = "Invalid date format. Showing today's report."
        else:
            specific_date = today
        start_date = specific_date
        end_date = specific_date
        orders_qs = Order.objects.filter(created_at__date=specific_date)

    elif filter_type == "weekly":
        start_date = today - datetime.timedelta(days=6)
        end_date = today
        orders_qs = Order.objects.filter(created_at__date__gte=start_date, created_at__date__lte=end_date)

    elif filter_type == "monthly":
        start_date = today.replace(day=1)
        end_date = today
        orders_qs = Order.objects.filter(created_at__date__gte=start_date, created_at__date__lte=end_date)

    elif filter_type == "yearly":
        start_date = today.replace(month=1, day=1)
        end_date = today
        orders_qs = Order.objects.filter(created_at__date__gte=start_date, created_at__date__lte=end_date)

    elif filter_type == "custom":
        start_date_str = request.GET.get("start_date", "").strip()
        end_date_str = request.GET.get("end_date", "").strip()
        orders_qs = Order.objects.all()

        if start_date_str:
            try:
                start_date = datetime.datetime.strptime(start_date_str, "%Y-%m-%d").date()
            except ValueError:
                date_error = "Invalid start date format (use YYYY-MM-DD)."
        if end_date_str:
            try:
                end_date = datetime.datetime.strptime(end_date_str, "%Y-%m-%d").date()
            except ValueError:
                date_error = "Invalid end date format (use YYYY-MM-DD)."

        if start_date and end_date and start_date > end_date:
            date_error = "Start date cannot be after end date."
            start_date, end_date = None, None
        else:
            if start_date:
                orders_qs = orders_qs.filter(created_at__date__gte=start_date)
            if end_date:
                orders_qs = orders_qs.filter(created_at__date__lte=end_date)
    else:
        filter_type = "all"
        orders_qs = Order.objects.all()

    if status_filter and status_filter != "all":
        orders_qs = orders_qs.filter(order_status=status_filter)

    orders = orders_qs.select_related("user").prefetch_related("items__product").order_by("-created_at")

    valid_sales_orders = orders_qs.exclude(order_status="CANCELLED")
    sales_count = valid_sales_orders.count()
    gross_order_amount = float(valid_sales_orders.aggregate(s=Sum("subtotal"))["s"] or 0.0)
    total_discount = float(valid_sales_orders.aggregate(s=Sum("discount"))["s"] or 0.0)
    net_sales = float(valid_sales_orders.aggregate(s=Sum("total_amount"))["s"] or 0.0)

    total_orders_count = orders_qs.count()
    cancelled_orders_count = orders_qs.filter(order_status="CANCELLED").count()
    delivered_orders_count = orders_qs.filter(order_status="DELIVERED").count()

    payment_stats = {
        "RAZORPAY": valid_sales_orders.filter(payment_method="RAZORPAY").count(),
        "COD": valid_sales_orders.filter(payment_method="COD").count(),
        "WALLET": valid_sales_orders.filter(payment_method="WALLET").count(),
    }
    total_valid = max(sales_count, 1)
    payment_percentages = {
        k: round((v / total_valid) * 100) for k, v in payment_stats.items()
    }

    chart_dict = {}
    for o in reversed(list(valid_sales_orders)):
        d_label = o.created_at.strftime("%d %b")
        chart_dict[d_label] = round(chart_dict.get(d_label, 0.0) + float(o.total_amount), 2)
    chart_labels = list(chart_dict.keys())
    chart_data = list(chart_dict.values())

    if export_format == "csv":
        response = HttpResponse(content_type="text/csv")
        response["Content-Disposition"] = f'attachment; filename="wegnex_sales_report_{timezone.now().strftime("%Y%m%d_%H%M%S")}.csv"'
        writer = csv.writer(response)
        writer.writerow(["WEGNEX LUXURY HOROLOGY - SALES REPORT"])
        writer.writerow([f"Filter: {filter_type.upper()}", f"Start Date: {start_date or 'All'}", f"End Date: {end_date or 'All'}", f"Generated: {timezone.now().strftime('%Y-%m-%d %H:%M')}"])
        writer.writerow([])
        writer.writerow(["SUMMARY METRICS"])
        writer.writerow(["Overall Sales Count", sales_count])
        writer.writerow(["Gross Order Amount (Rs.)", f"{gross_order_amount:.2f}"])
        writer.writerow(["Coupon / Discount Deductions (Rs.)", f"{total_discount:.2f}"])
        writer.writerow(["Net Sales Revenue (Rs.)", f"{net_sales:.2f}"])
        writer.writerow(["Total Orders (incl. cancelled)", total_orders_count])
        writer.writerow(["Cancelled Orders", cancelled_orders_count])
        writer.writerow([])
        writer.writerow(["Order Number", "Date & Time", "Customer Name", "Customer Email", "Customer Phone", "Items Summary", "Gross Amount (Rs.)", "Coupon Code", "Discount (Rs.)", "Net Total (Rs.)", "Payment Method", "Payment Status", "Order Status"])
        for o in orders:
            items_str = "; ".join([f"{item.product_name} x{item.quantity}" for item in o.items.all()]) or "N/A"
            writer.writerow([
                o.order_number,
                o.created_at.strftime("%Y-%m-%d %H:%M"),
                o.shipping_full_name or o.user.get_full_name() or o.user.username,
                o.user.email,
                o.shipping_phone or "N/A",
                items_str,
                f"{o.subtotal:.2f}",
                o.coupon_code or "None",
                f"{o.discount:.2f}",
                f"{o.total_amount:.2f}",
                o.payment_method,
                o.payment_status,
                o.order_status,
            ])
        return response

    if export_format == "excel":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        ws = wb.active
        ws.title = "Sales Report"

        title_font = Font(name="Arial", size=13, bold=True, color="FFFFFF")
        header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        bold_font = Font(name="Arial", size=10, bold=True)
        normal_font = Font(name="Arial", size=9)

        dark_fill = PatternFill(start_color="111111", end_color="111111", fill_type="solid")
        header_fill = PatternFill(start_color="232323", end_color="232323", fill_type="solid")
        summary_head_fill = PatternFill(start_color="EAEAEA", end_color="EAEAEA", fill_type="solid")

        ws.merge_cells("A1:L1")
        ws["A1"] = "WEGNEX LUXURY TIMEPIECES — SALES REPORT"
        ws["A1"].font = title_font
        ws["A1"].fill = dark_fill
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 34

        ws.append([f"Filter: {filter_type.upper()}", f"Generated: {timezone.now().strftime('%Y-%m-%d %H:%M')}", "", "", "", "", "", "", "", "", "", ""])
        ws.append([])

        ws.append(["SUMMARY METRICS", "", "", ""])
        ws.append(["Overall Sales Count", sales_count, "Gross Order Amount (Rs.)", gross_order_amount])
        ws.append(["Total Orders Placed", total_orders_count, "Coupon/Discount Deductions (Rs.)", total_discount])
        ws.append(["Cancelled Orders", cancelled_orders_count, "Net Sales Revenue (Rs.)", net_sales])
        ws.append([])

        for r_idx in range(4, 8):
            for c_idx in range(1, 5):
                cell = ws.cell(row=r_idx, column=c_idx)
                if c_idx in (1, 3):
                    cell.font = bold_font
                    cell.fill = summary_head_fill
                else:
                    cell.font = normal_font

        headers = ["Order #", "Date & Time", "Customer Name", "Customer Email", "Items Summary", "Gross Amount (Rs.)", "Coupon Code", "Discount (Rs.)", "Net Total (Rs.)", "Payment Method", "Payment Status", "Order Status"]
        ws.append(headers)
        h_row = ws.max_row
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=h_row, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[h_row].height = 24

        for o in orders:
            items_str = ", ".join([f"{item.product_name} x{item.quantity}" for item in o.items.all()]) or "N/A"
            ws.append([
                o.order_number,
                o.created_at.strftime("%Y-%m-%d %H:%M"),
                o.shipping_full_name or o.user.get_full_name() or o.user.username,
                o.user.email,
                items_str,
                float(o.subtotal),
                o.coupon_code or "—",
                float(o.discount),
                float(o.total_amount),
                o.payment_method,
                o.payment_status,
                o.order_status,
            ])
            cur_row = ws.max_row
            ws.cell(row=cur_row, column=6).number_format = "#,##0.00"
            ws.cell(row=cur_row, column=8).number_format = "#,##0.00"
            ws.cell(row=cur_row, column=9).number_format = "#,##0.00"

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = f'attachment; filename="wegnex_sales_report_{timezone.now().strftime("%Y%m%d_%H%M%S")}.xlsx"'
        return response

    if export_format == "pdf":
        from reportlab.lib.pagesizes import letter, landscape
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib import colors

        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=landscape(letter), leftMargin=24, rightMargin=24, topMargin=24, bottomMargin=24)
        styles = getSampleStyleSheet()
        elements = []

        title_style = ParagraphStyle("RepTitle", parent=styles["Heading1"], fontSize=15, fontName="Helvetica-Bold", textColor=colors.HexColor("#111111"), spaceAfter=3)
        subtitle_style = ParagraphStyle("RepSubtitle", parent=styles["Normal"], fontSize=8, fontName="Helvetica", textColor=colors.HexColor("#666666"), spaceAfter=12)
        cell_style = ParagraphStyle("TCell", parent=styles["Normal"], fontSize=7, leading=9)
        cell_bold = ParagraphStyle("TCellB", parent=styles["Normal"], fontSize=7, leading=9, fontName="Helvetica-Bold")

        elements.append(Paragraph("WEGNEX LUXURY HOROLOGY — SALES REPORT", title_style))
        elements.append(Paragraph(f"Period: <b>{filter_type.upper()}</b> | Range: {start_date or 'Beginning'} to {end_date or 'Present'} | Generated: {timezone.now().strftime('%Y-%m-%d %H:%M')}", subtitle_style))

        sum_data = [
            ["Overall Sales Count", "Gross Order Amount", "Discounts / Coupon", "Net Sales Revenue"],
            [str(sales_count), f"Rs. {gross_order_amount:,.2f}", f"-Rs. {total_discount:,.2f}", f"Rs. {net_sales:,.2f}"]
        ]
        sum_table = Table(sum_data, colWidths=[185, 185, 185, 185])
        sum_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1a1a1a")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), 9),
            ("BACKGROUND", (0, 1), (-1, 1), colors.HexColor("#f5f5f5")),
            ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"),
            ("FONTSIZE", (0, 1), (-1, 1), 10),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cccccc")),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
        ]))
        elements.append(sum_table)
        elements.append(Spacer(1, 12))

        t_data = [["Order #", "Date", "Customer", "Gross (Rs.)", "Discount", "Net Total (Rs.)", "Payment", "Status"]]
        for o in orders:
            cust = (o.shipping_full_name or o.user.username)[:16]
            disc_text = f"Rs. {o.discount:.2f}"
            if o.coupon_code:
                disc_text += f" ({o.coupon_code})"
            t_data.append([
                Paragraph(o.order_number, cell_style),
                Paragraph(o.created_at.strftime("%Y-%m-%d"), cell_style),
                Paragraph(cust, cell_style),
                Paragraph(f"{o.subtotal:,.2f}", cell_style),
                Paragraph(disc_text, cell_style),
                Paragraph(f"{o.total_amount:,.2f}", cell_bold),
                Paragraph(o.payment_method, cell_style),
                Paragraph(o.order_status, cell_style),
            ])

        table = Table(t_data, colWidths=[120, 75, 120, 85, 95, 95, 75, 85])
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#222222")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), 8),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e0e0e0")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#fafafa")]),
        ]))
        elements.append(table)
        doc.build(elements)
        buf.seek(0)
        response = HttpResponse(buf.getvalue(), content_type="application/pdf")
        response["Content-Disposition"] = f'attachment; filename="wegnex_sales_report_{timezone.now().strftime("%Y%m%d_%H%M%S")}.pdf"'
        return response

    return render(request, "sales_report.html", {
        "admin_user": request.user,
        "orders": orders,
        "filter_type": filter_type,
        "status_filter": status_filter,
        "start_date": start_date.strftime("%Y-%m-%d") if start_date else "",
        "end_date": end_date.strftime("%Y-%m-%d") if end_date else "",
        "specific_date": specific_date.strftime("%Y-%m-%d") if specific_date else "",
        "date_error": date_error,
        "sales_count": sales_count,
        "gross_order_amount": gross_order_amount,
        "total_discount": total_discount,
        "net_sales": net_sales,
        "total_orders_count": total_orders_count,
        "cancelled_orders_count": cancelled_orders_count,
        "delivered_orders_count": delivered_orders_count,
        "payment_stats": payment_stats,
        "payment_percentages": payment_percentages,
        "chart_labels": chart_labels,
        "chart_data": chart_data,
    })